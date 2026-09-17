"""Application runtime Excel import/export (no external service required)."""
import base64
from datetime import date, datetime
import io
import re
import zipfile
from urllib.parse import quote
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import xlrd

MAX_ROWS = 1000

def normalize(value):
    return re.sub(r'[\s*（）()：:、/\-_]+', '', str(value or '')).lower()

def read_file(payload):
    filename = str(payload.get('filename', '')).lower()
    if not filename.endswith(('.xlsx', '.xls')):
        raise ValueError('请选择 .xlsx 或 .xls Excel 文件')
    try:
        raw = base64.b64decode(payload.get('file', ''), validate=True)
        if not raw or len(raw) > 8 * 1024 * 1024:
            raise ValueError('Excel文件大小须在8MB以内')
        sheets = {}
        if filename.endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('Excel解压后内容过大，请拆分文件')
            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                names = workbook.sheetnames
                name = payload.get('sheet') or names[0]
                if name not in names:
                    raise ValueError('工作表不存在')
                sheet = workbook[name]
                if sheet.max_column > 150 or sheet.max_row > MAX_ROWS + 30:
                    raise ValueError('每次最多导入1000条数据、150列；请拆分工作表')
                rows = [[(cell.value, cell.data_type == 'f') for cell in row] for row in sheet.iter_rows()]
            finally:
                workbook.close()
        else:
            workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
            try:
                names = workbook.sheet_names()
                name = payload.get('sheet') or names[0]
                if name not in names:
                    raise ValueError('工作表不存在')
                sheet = workbook.sheet_by_name(name)
                if sheet.ncols > 150 or sheet.nrows > MAX_ROWS + 30:
                    raise ValueError('每次最多导入1000条数据、150列；请拆分工作表')
                rows = []
                for row_index in range(sheet.nrows):
                    row = []
                    for cell in sheet.row(row_index):
                        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                        row.append((value, False))
                    rows.append(row)
            finally:
                workbook.release_resources()
        return names, name, rows
    except ValueError:
        raise
    except Exception:
        raise ValueError('无法读取Excel文件，请检查文件格式，取消密码保护后重试')

def column_mapping(fields, rows, custom=None):
    aliases = {'name':['学校名称','院校名称','学习名称','学生姓名','应聘者姓名'], 'province':['省份','省'], 'city':['城市','市'], 'school_name':['毕业学校','学校'], 'phone':['手机号','联系电话','手机'], 'education':['最高学历'], 'position':['岗位','意向岗位','应聘岗位'], 'candidate_id':['应聘人','入职人员','应聘人员','人员姓名'], 'school_id':['关联院校','院校档案'], 'level':['院校层次'], 'employment_account':['就业网用户名']}
    aliases.update(major=['所学专业'], graduation=['毕业日期','毕业年月'], phone=['手机号','联系电话','手机','手机号码','电话号码'], email=['邮箱'], position=['岗位','意向岗位','应聘岗位'])
    lookup = {}
    for field in fields:
        if field.get('readonly') or field.get('hidden') or field.get('type') == 'attachment':
            continue
        for label in [field['label'], field['key'], *aliases.get(field['key'], [])]:
            lookup[normalize(label)] = field['key']
        if field.get('ref'):
            lookup[normalize(field['label'] + '编号')] = field['key']
    if not rows:
        raise ValueError('工作表为空')
    header_index = max(range(min(20, len(rows))), key=lambda index: sum(normalize(value) in lookup for value, _ in rows[index]))
    headers = [str(value or '') for value, _ in rows[header_index]]
    mapping = [lookup.get(normalize(header), '') for header in headers]
    # Exported reference IDs take precedence over their human-readable name column.
    for index, header in enumerate(headers):
        if header.endswith('编号') and mapping[index]:
            for other in range(len(mapping)):
                if other != index and mapping[other] == mapping[index]:
                    mapping[other] = ''
    if custom is not None:
        allowed = {field['key'] for field in fields if not field.get('readonly') and not field.get('hidden') and field.get('type') != 'attachment'}
        if not isinstance(custom, list) or len(custom) != len(headers) or any(key and key not in allowed for key in custom):
            raise ValueError('列映射无效，请重新识别')
        mapping = custom
    keys = [key for key in mapping if key]
    if len(keys) != len(set(keys)):
        raise ValueError('多个Excel列对应同一字段，请调整列名或列映射')
    return header_index, headers, mapping

def prepare_rows(entity, spec, payload, conn, validate):
    names, name, rows = read_file(payload)
    if not rows:
        return {'sheets': names, 'sheet': name, 'headers': [], 'mapping': [], 'header_row': 1, 'rows': [], 'errors': [{'line': 1, 'error': '工作表为空，请选择其他工作表'}]}
    header, headers, mapping = column_mapping(spec['fields'], rows, payload.get('mapping'))
    if not any(mapping):
        return {'sheets': names, 'sheet': name, 'headers': headers, 'mapping': mapping, 'header_row': header + 1, 'rows': [], 'errors': [{'line': header + 1, 'error': '未识别到匹配字段，请调整下方列对应关系或选择其他工作表'}]}
    fields = {field['key']: field for field in spec['fields']}
    data, errors = [], []
    for line, cells in enumerate(rows[header + 1:], header + 2):
        if not any(value is not None and str(value).strip() for value, _ in cells):
            continue
        try:
            values = {}
            for index, key in enumerate(mapping):
                if not key:
                    continue
                value, formula = cells[index] if index < len(cells) else (None, False)
                if formula:
                    raise ValueError('含公式，请在Excel中粘贴为值后导入')
                if isinstance(value, (datetime, date)):
                    value = value.strftime('%Y-%m-%dT%H:%M:%S') if fields[key].get('type') == 'datetime-local' else value.strftime('%Y-%m-%d')
                elif isinstance(value, bool):
                    value = '是' if value else '否'
                elif isinstance(value, float) and value.is_integer():
                    value = str(int(value))
                value = str(value) if value is not None else ''
                field = fields[key]
                if field.get('ref') and value:
                    if not value.isdigit() or (field['ref'] == 'candidates' and len(value) == 11):
                        ref = field['ref']
                        matches = conn.execute(f'SELECT id FROM {ref} WHERE name=?' + (' OR phone=?' if ref == 'candidates' else ''), (value, value) if ref == 'candidates' else (value,)).fetchall()
                        if len(matches) != 1:
                            raise ValueError(field['label'] + '无法唯一匹配，请填写档案编号或唯一手机号/名称')
                        value = str(matches[0]['id'])
                values[key] = value
            # Omitted stage columns receive the same defaults as the create form.
            if 'status' in fields and 'status' not in values:
                values['status'] = fields['status']['options'][0]
            data.append({'line': line, 'values': validate(entity, values, conn)})
        except ValueError as error:
            errors.append({'line': line, 'error': str(error)})
    if len(data) + len(errors) > MAX_ROWS:
        raise ValueError('每次最多导入1000条，请拆分文件')
    if not data and not errors:
        errors.append({'line': header + 1, 'error': '表头下方没有可导入的数据，请选择其他工作表'})
    return {'sheets': names, 'sheet': name, 'headers': headers, 'mapping': mapping, 'header_row': header + 1, 'rows': data, 'errors': errors}

def workbook_bytes(spec, rows, template=False, attachment_base_url=''):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = spec['short']
    fields = [field for field in spec['fields'] if not field.get('hidden') and (not field.get('readonly') and field.get('type') != 'attachment' if template else field.get('export', True))]
    columns = [] if template else [('编号', 'id', None)]
    for field in fields:
        if field.get('ref') and not template:
            columns.append((field['label'] + '编号', field['key'], None))
            columns.append((field['label'], field['key'] + '_label', None))
        else:
            columns.append((field['label'], field['key'], field.get('type')))
    if not template and not any(field['key'] == 'created_date' for field in fields):
        columns.append(('创建时间', 'created_at', None))
    sheet.append([label for label, _, _ in columns])
    for row_index, record in enumerate(rows, 2):
        for column_index, (_, key, kind) in enumerate(columns, 1):
            value = record.get(key + '_name', '') if kind == 'attachment' else record.get(key, '')
            if value and kind in ('number', 'integer'):
                value = float(value) if kind == 'number' else int(value)
            elif value and kind == 'date':
                value = datetime.strptime(value, '%Y-%m-%d')
            elif value and kind == 'month':
                value = datetime.strptime(value, '%Y-%m')
            elif value and kind == 'datetime-local':
                value = datetime.fromisoformat(value)
            elif value and key == 'created_date':
                value = datetime.strptime(value, '%Y-%m-%d %H:%M:%S')
            cell = sheet.cell(row_index, column_index, value)
            if kind == 'attachment' and value and record.get(key) and attachment_base_url:
                cell.hyperlink = attachment_base_url.rstrip('/') + '/api/attachments/' + quote(str(record[key]), safe='') + '?preview=1'
                cell.style = 'Hyperlink'
            if isinstance(value, str):
                cell.data_type = 's'  # Literal text, never executable formula content.
            if kind == 'date':
                cell.number_format = 'yyyy-mm-dd'
            elif kind == 'month':
                cell.number_format = 'yyyy-mm'
            elif key == 'interview_at':
                cell.number_format = 'yyyy-mm-dd hh:mm'
            elif key == 'created_date' or kind == 'datetime-local':
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='2159D0')
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 28
    for index, (label, key, _) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24 if key in ('created_date', 'created_at', 'applied_at', 'interview_at') else max(16, min(36, len(label) * 2 + 4))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
