import base64
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.request import build_opener, HTTPCookieProcessor, Request
from urllib.error import HTTPError
import http.cookiejar
import server
import openpyxl
from excel_io import workbook_bytes


class TableToolsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.original=server.DATA
        server.DATA=Path(cls.temp.name)
        server.initialize()
        cls.http=server.ThreadingHTTPServer(('127.0.0.1',0),server.Handler)
        cls.thread=threading.Thread(target=cls.http.serve_forever,daemon=True)
        cls.thread.start()
        cls.url='http://127.0.0.1:'+str(cls.http.server_port)
        cls.client=build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))
        password=(server.DATA/'initial-password.txt').read_text()
        cls.call('login',{'username':'admin','password':password})

    @classmethod
    def tearDownClass(cls):
        cls.http.shutdown(); cls.http.server_close(); cls.thread.join()
        server.DATA=cls.original
        server.IMPORTS.clear()
        cls.temp.cleanup()

    @classmethod
    def call(cls,path,data=None,public=False,method=None):
        request=Request(cls.url+'/api/'+path,data=json.dumps(data).encode() if data is not None else None,headers={'Content-Type':'application/json'},method=method)
        try:
            response=(build_opener() if public else cls.client).open(request)
        except HTTPError as error:
            response=error
        with response:
            raw=response.read()
            return response.status,json.loads(raw) if response.headers.get_content_type()=='application/json' else raw

    @staticmethod
    def file(headers,rows):
        book=openpyxl.Workbook(); sheet=book.active
        sheet.append(headers)
        for row in rows: sheet.append(row)
        stream=io.BytesIO();book.save(stream);book.close()
        return {'filename':'test.xlsx','file':base64.b64encode(stream.getvalue()).decode()}

    def test_01_preview_atomic_import_and_replay(self):
        payload=self.file(['学校名称','省份','城市'],[['甲大学','江苏','南京市'],['乙大学','湖北省','武汉市']])
        status,preview=self.call('schools/import-preview',payload)
        self.assertEqual(status,200)
        self.assertEqual(preview['total'],2)
        self.assertTrue(preview['token'])
        self.assertEqual(self.call('schools')[1],[])
        status,result=self.call('schools/import-commit',{'token':preview['token']})
        self.assertEqual(status,201)
        self.assertEqual(result['count'],2)
        self.assertEqual(self.call('schools/import-commit',{'token':preview['token']})[0],400)
        self.assertEqual(len(self.call('schools')[1]),2)
        bad=self.file(['学校名称','所在省','所在市'],[['可用大学','江苏省','南京市'],['错误大学','江苏省','武汉市']])
        preview=self.call('schools/import-preview',bad)[1]
        self.assertIsNone(preview['token'])
        self.assertEqual(preview['errors'][0]['line'],3)
        self.assertEqual(len(self.call('schools')[1]),2)

    def test_02_all_export_templates_and_qr(self):
        for entity,spec in server.SCHEMAS.items():
            for action in ['export.xlsx','template.xlsx']:
                status,content=self.call(entity+'/'+action)
                self.assertEqual(status,200)
                book=openpyxl.load_workbook(io.BytesIO(content))
                self.assertEqual(book.active.freeze_panes,'A2')
                headers=[cell.value for cell in book.active[1]]
                first_field=next(field for field in spec['fields'] if not field.get('readonly') and not field.get('hidden'))
                self.assertIn(first_field['label'],headers)
                if action=='export.xlsx':self.assertNotIn('就业网密码',headers)
                book.close()
        status,result=self.call('schools/qr-print',{'ids':[1,2]})
        self.assertEqual(status,200)
        self.assertEqual(len(result['items']),2)
        for item in result['items']:
            self.assertIn('module=schools&record='+str(item['id']),item['url'])
            self.assertIn(b'<svg',base64.b64decode(item['image'].split(',')[1]))
        self.assertEqual(self.call('schools/qr-print',{'ids':[9999]})[0],400)
        self.assertEqual(self.call('schools/qr-print',{'ids':[1]},public=True)[0],401)
        self.assertEqual(self.call('schools/export.xlsx',public=True)[0],401)

        school_ids=[row['id'] for row in self.call('schools')[1]]
        selected=self.call('schools/export.xlsx',{'ids':[school_ids[0]]},method='POST')[1]
        selected_book=openpyxl.load_workbook(io.BytesIO(selected));selected_sheet=selected_book.active
        self.assertEqual(selected_sheet.max_row,2)
        self.assertEqual(selected_sheet['A2'].value,school_ids[0])
        selected_book.close()
        self.assertEqual(self.call('schools/export.xlsx',{'ids':[]},method='POST')[0],400)
        self.assertEqual(self.call('schools/export.xlsx',{'ids':[999999]},method='POST')[0],400)

    def test_03_formulas_types_mapping_and_duplicates(self):
        formula=self.file(['学校名称','所在省'],[['=1+1','江苏省']])
        self.assertIn('公式',self.call('schools/import-preview',formula)[1]['errors'][0]['error'])
        exported=workbook_bytes(server.SCHEMAS['questions'],[{'id':1,'question':'=1+1'}])
        book=openpyxl.load_workbook(io.BytesIO(exported))
        self.assertEqual(book.active['B2'].value,'=1+1')
        self.assertEqual(book.active['B2'].data_type,'s')
        book.close()
        row=['王同学',13800138000,'甲大学','工程','本科','技术']
        payload=self.file(['姓名','手机号','毕业院校','所学专业','学历','意向岗位'],[row,row])
        preview=self.call('candidates/import-preview',payload)[1]
        self.assertIsNone(preview['token'])
        self.assertEqual(self.call('candidates')[1],[])
        payload=self.file(['姓名','手机号','毕业院校','所学专业','学历','意向岗位'],[row])
        preview=self.call('candidates/import-preview',payload)[1]
        self.assertTrue(preview['token'])
        self.assertEqual(self.call('candidates/import-commit',{'token':preview['token']})[0],201)
        payload=self.file(['姓名','毕业院校','联系电话','面试时间','流程状态'],[['王同学','甲大学','13800138000','2026-09-07T10:30','草稿']])
        preview=self.call('interviews/import-preview',payload)[1]
        self.assertTrue(preview['token'])
        self.assertEqual(self.call('interviews/import-commit',{'token':preview['token']})[0],201)
        self.assertEqual(self.call('interviews')[1][0]['name'],'王同学')

    def test_04_commit_race_rolls_back_entire_batch(self):
        rows=[['先成功',13900139001,'甲大学','工程','本科','技术'],['后冲突',13900139002,'甲大学','工程','本科','技术']]
        payload=self.file(['姓名','手机号','毕业院校','所学专业','学历','意向岗位'],rows)
        preview=self.call('candidates/import-preview',payload)[1]
        self.assertTrue(preview['token'])
        person={'name':'抢先登记','phone':'13900139002','school_name':'甲大学','major':'工程','education':'本科','position':'技术','status':'待筛选'}
        self.call('candidates',person)
        before=len(self.call('candidates')[1])
        self.assertEqual(self.call('candidates/import-commit',{'token':preview['token']})[0],400)
        self.assertEqual(len(self.call('candidates')[1]),before)

    def test_05_remaining_modules_and_sheet_mapping(self):
        examples={
            'applications': (['联系电话','日期','岗位'],[['13800138000','2026-09-07','工程技术']]),
            'employees': (['姓名','岗位','身高（cm）','体重（kg）','招聘途径'],[['王同学','技术员','178','70','校园线下']]),
            'questions': (['问题','通用回答话术'],[['需要携带哪些材料？','请携带身份证和毕业证。']])
        }
        for entity,(headers,rows) in examples.items():
            preview=self.call(entity+'/import-preview',self.file(headers,rows))[1]
            self.assertTrue(preview['token'],preview)
            self.assertEqual(self.call(entity+'/import-commit',{'token':preview['token']})[0],201)
            self.assertEqual(self.call(entity+'/qr-print',{'ids':[1]})[0],200)
        book=openpyxl.Workbook();book.active.title='说明';book.active.append(['先阅读说明'])
        sheet=book.create_sheet('数据');sheet.append(['自定义校名','省份']);sheet.append(['映射大学','江苏省'])
        stream=io.BytesIO();book.save(stream);book.close()
        payload={'filename':'multiple.xlsx','file':base64.b64encode(stream.getvalue()).decode()}
        preview=self.call('schools/import-preview',payload)[1]
        self.assertIn('数据',preview['sheets'])
        self.assertIsNone(preview['token'])
        preview=self.call('schools/import-preview',dict(payload,sheet='数据',mapping=['name','province']))[1]
        self.assertTrue(preview['token'])
        self.assertEqual(self.call('schools/import-commit',{'token':preview['token']})[0],201)

    def test_06_candidate_dates_resume_and_export(self):
        person={'name':'新档案','phone':'13700137000','position':'技术员','birthdate':'2000-10-01','channel':'校园平台','created_at':'1999-01-01','created_date':'1999-01-01','age':99,'cohort':'错误届别','education_experiences':[{'education':'专科','school_name':'甲大学','major':'机械','enrollment':'2023-09','graduation':'2026-7'}]}
        person['_resume_upload']={'name':'简历.pdf','content':base64.b64encode(b'%PDF-1.4\nresume-test').decode()}
        status,created=self.call('candidates',person)
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertRegex(row['created_date'],r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$')
        self.assertNotIn('1999',row['created_date'])
        self.assertEqual(row['age'],'99')
        self.assertEqual(row['graduation'],'2026-07')
        self.assertEqual(row['cohort'],'2026届')
        self.assertEqual(row['resume_name'],'简历.pdf')
        old_resume=row['resume']
        self.assertEqual(self.call('attachments/'+old_resume)[1],b'%PDF-1.4\nresume-test')
        self.assertEqual(self.call('attachments/'+old_resume,public=True)[0],401)
        status,_=self.call('candidates/'+str(created['id']),{'birthdate':'2001-01-01','education_experiences':[{'education':'专科','school_name':'甲大学','major':'机械','enrollment':'2023-09','graduation':'2027-06'}],'created_date':'2000-01-01 00:00:00','age':41},method='PUT')
        self.assertEqual(status,200)
        updated=next(item for item in self.call('candidates')[1] if item['id']==created['id'])
        self.assertEqual(updated['created_date'],row['created_date'])
        self.assertEqual(updated['cohort'],'2027届')
        self.assertEqual(updated['age'],'41')
        self.assertEqual(updated['resume'],old_resume)
        self.assertEqual(self.call('candidates/'+str(created['id']),{'birthdate':'2999-01-01'},method='PUT')[0],400)
        self.assertEqual(self.call('candidates/'+str(created['id']),{'education_experiences':[{'education':'专科','school_name':'甲大学','major':'机械','graduation':'2026-13'}]},method='PUT')[0],400)
        exported=self.call('candidates/export.xlsx',{'ids':[created['id']]},method='POST')[1]
        workbook=openpyxl.load_workbook(io.BytesIO(exported));sheet=workbook.active
        headers=[cell.value for cell in sheet[1]]
        matching=next(values for values in sheet.iter_rows(min_row=2) if values[0].value==created['id'])
        self.assertEqual(matching[headers.index('创建日期')].number_format,'yyyy-mm-dd hh:mm:ss')
        self.assertEqual(matching[headers.index('毕业时间')].number_format,'yyyy-mm')
        resume_cell=matching[headers.index('简历附件')]
        self.assertEqual(resume_cell.value,'简历.pdf')
        self.assertIsNotNone(resume_cell.hyperlink)
        self.assertTrue(resume_cell.hyperlink.target.endswith('/api/attachments/'+old_resume+'?preview=1'))
        workbook.close()
        self.assertEqual(self.call('candidates/'+str(created['id']),{'_remove_resume':True},method='PUT')[0],200)
        self.assertEqual(self.call('attachments/'+old_resume)[0],404)
        public=dict(person,phone='13600136000')
        self.assertEqual(self.call('public/apply',public,public=True)[0],201)
        with server.db() as conn:
            before=conn.execute('SELECT COUNT(*) FROM attachments').fetchone()[0]
        self.assertEqual(self.call('public/apply',public,public=True)[0],200)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM attachments').fetchone()[0],before)
        self.assertEqual(self.call('candidates',dict(person,phone='13500135000',_resume_upload={'name':'resume.pdf','content':base64.b64encode(b'not-pdf').decode()}))[0],400)

    def test_07_optional_position_and_application_time(self):
        from datetime import datetime
        person={'name':'无岗位人才','phone':'13400134000','school_name':'甲大学','major':'工程','education':'本科','applied_at':'2026-09-07T10:25:36'}
        self.call('candidates',person)
        status,result=self.call('public/apply',person,public=True)
        self.assertEqual(status,201)
        row=next(item for item in self.call('candidates')[1] if item['id']==result['candidate_id'])
        self.assertEqual(row['position'],'')
        self.assertEqual(row['applied_at'],'2026-09-07T10:25:36')
        application=next(item for item in self.call('applications')[1] if item['candidate_id']==result['candidate_id'])
        self.assertEqual(application['position'],'')
        self.assertEqual(self.call('candidates/'+str(result['candidate_id']),{'applied_at':''},method='PUT')[0],200)
        self.assertEqual(self.call('candidates',dict(person,phone='13300133000',applied_at='2026-02-30T09:00'))[0],400)
        payload=self.file(['姓名','联系方式','毕业院校','专业','学历','投递时间'],[['导入无岗位','13300133000','甲大学','工程','本科',datetime(2026,9,7,11,22,33)]])
        preview=self.call('candidates/import-preview',payload)[1]
        self.assertTrue(preview['token'],preview)
        status,result=self.call('candidates/import-commit',{'token':preview['token']})
        self.assertEqual(status,201)
        row=next(item for item in self.call('candidates')[1] if item['id']==result['ids'][0])
        self.assertEqual(row['applied_at'],'2026-09-07T11:22:33')
        self.assertEqual(row['position'],'')
        workbook=openpyxl.load_workbook(io.BytesIO(self.call('candidates/export.xlsx')[1]))
        sheet=workbook.active;headers=[cell.value for cell in sheet[1]]
        values=next(cells for cells in sheet.iter_rows(min_row=2) if cells[0].value==row['id'])
        self.assertEqual(values[headers.index('投递时间')].value,datetime(2026,9,7,11,22,33))
        workbook.close()

    def test_08_batch_delete_atomicity_and_all_modules(self):
        person={'name':'批量删除测试','phone':'13200132000','school_name':'甲大学','major':'工程','education':'本科','_resume_upload':{'name':'resume.pdf','content':base64.b64encode(b'%PDF-1.4\ntest').decode()}}
        status,created=self.call('candidates',person)
        self.assertEqual(status,201)
        free_id=created['id']
        row=next(item for item in self.call('candidates')[1] if item['id']==free_id)
        # Candidate 1 is referenced by earlier interview/application records.
        preview=self.call('candidates/delete-preview',{'ids':[free_id,1]})
        self.assertEqual(preview[0],200)
        self.assertGreater(preview[1]['count'],2)
        self.assertTrue(any(item['id']==free_id for item in self.call('candidates')[1]))
        self.assertEqual(self.call('attachments/'+row['resume'])[0],200)
        self.assertEqual(self.call('candidates/batch-delete',{'ids':[free_id,999999]})[0],400)
        self.assertEqual(self.call('candidates/batch-delete',{'ids':[free_id]},public=True)[0],401)
        status,result=self.call('candidates/batch-delete',{'ids':[free_id,free_id]})
        self.assertEqual((status,result['count']),(200,1))
        self.assertEqual(self.call('attachments/'+row['resume'])[0],404)
        for invalid in [[],[True],['1'],[0],list(range(1,1002))]:
            self.assertEqual(self.call('schools/batch-delete',{'ids':invalid})[0],400)
        examples={
            'schools':{'name':'删除院校','province':'江苏省'},
            'applications':{'candidate_id':1,'date':'2026-09-07','position':'测试','status':'草稿'},
            'interviews':{'candidate_id':1,'name':'王同学','school_name':'甲大学','date':'2026-09-07','interviewer':'测试','status':'草稿'},
            'questions':{'title':'删除问题','category':'其他','question':'测试','status':'待答复'}
        }
        for entity,payload in examples.items():
            ids=[self.call(entity,payload)[1]['id'] for _ in range(2)]
            status,result=self.call(entity+'/batch-delete',{'ids':ids})
            self.assertEqual((status,result['count']),(200,2))
            self.assertFalse(set(ids)&{row['id'] for row in self.call(entity)[1]})
        employee_ids=[row['id'] for row in self.call('employees')[1]]
        self.assertTrue(employee_ids)
        self.assertEqual(self.call('employees/batch-delete',{'ids':employee_ids})[1]['count'],len(employee_ids))

    def test_09_application_profile_and_multiple_attachments(self):
        school=self.call('schools',{'name':'应聘登记测试大学','province':'湖北省','is_985':'否','is_211':'是','is_double_first':'是'})[1]['id']
        candidate=self.call('candidates',{'name':'应聘测试','phone':'13100131000','school_id':school,'school_name':'应聘登记测试大学','major':'自动化','education':'本科','graduation':'2026-07','birthdate':'2003-06-10','channel':'校园平台','_resume_upload':{'name':'人才简历.pdf','content':base64.b64encode(b'%PDF-1.4\nresume').decode()}})[1]['id']
        self.assertEqual(self.call('application-prefill?id='+str(candidate))[1]['is_211'],'是')
        pdf={'name':'附件.pdf','content':base64.b64encode(b'%PDF-1.4\nattachment').decode()}
        png={'name':'签名.png','content':'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1kAAAAASUVORK5CYII='}
        values={'candidate_id':candidate,'screening':'合格','needs_first_interview':'是','ethnicity':'汉族','enrollment_type':'统招专升本','degree':'学士','enrollment':'2022-09','language':'英语','english_level':'CET-5','other_language':'日语','major_category':'电气类','other_major':'控制工程','political_status':'共青团员','identity_card':'110101200306100017','health_condition':'否','language_scores':'雅思 6.5','position':'技术岗','date':'2026-09-07','is_elite':'是','_transcript_upload':pdf,'_certificates_upload':pdf,'_signature_upload':png,'created_date':'伪造','age':99,'cohort':'错误'}
        status,created=self.call('applications',values)
        self.assertEqual(status,201)
        row=next(item for item in self.call('applications')[1] if item['id']==created['id'])
        self.assertEqual(row['name'],'应聘测试')
        self.assertEqual(row['school_name'],'应聘登记测试大学')
        self.assertEqual(row['english_level'],'CET-5')
        self.assertEqual(row['cohort'],'2026届')
        self.assertEqual(row['resume_name'],'人才简历.pdf')
        self.assertNotEqual(row['created_date'],'伪造')
        for key in ('resume','transcript','certificates','signature'):
            self.assertTrue(row[key])
            self.assertEqual(self.call('attachments/'+row[key])[0],200)
            self.assertEqual(self.call('attachments/'+row[key],public=True)[0],401)
        for updates in ({'screening':'未知'},{'english_level':'CET-6'},{'channel':'网络投递'}):
            self.assertEqual(self.call('applications/'+str(created['id']),updates,method='PUT')[0],400)
        self.assertEqual(self.call('applications/'+str(created['id']),{'_remove_transcript':True},method='PUT')[0],200)
        self.assertEqual(self.call('attachments/'+row['transcript'])[0],404)
        exported=self.call('applications/export.xlsx')[1]
        workbook=openpyxl.load_workbook(io.BytesIO(exported));sheet=workbook.active
        headers=[cell.value for cell in sheet[1]]
        values=next(cells for cells in sheet.iter_rows(min_row=2) if cells[0].value==row['id'])
        self.assertEqual(values[headers.index('证书附件')].value,'附件.pdf')
        self.assertEqual(values[headers.index('本人签名')].value,'签名.png')
        workbook.close()
        self.assertEqual(self.call('applications/batch-delete',{'ids':[created['id']]})[0],200)
        self.assertEqual(self.call('attachments/'+row['certificates'])[0],404)
        self.assertEqual(self.call('attachments/'+row['signature'])[0],404)
        self.assertEqual(self.call('attachments/'+row['resume'])[0],200)
        # Standalone application import does not require an existing talent record.
        payload=self.file(['姓名','毕业院校','学历','初筛情况','是否需要一面','毕业时间','简历渠道'],[['独立登记','测试大学','硕士','合格','否','2027-06','智联招聘']])
        preview=self.call('applications/import-preview',payload)[1]
        self.assertTrue(preview['token'],preview)
        self.assertEqual(self.call('applications/import-commit',{'token':preview['token']})[0],201)

    def test_10_application_phone_lookup(self):
        result=self.call('application-prefill?phone=13100131000')[1]
        self.assertTrue(result['matched'])
        self.assertEqual(result['values']['name'],'应聘测试')
        self.assertFalse(self.call('application-prefill?phone=12900000000')[1]['matched'])
        self.assertEqual(self.call('application-prefill?phone=13100131000',public=True)[0],401)
        status,created=self.call('applications',{'phone':'13100131000'})
        self.assertEqual(status,201)
        row=next(item for item in self.call('applications')[1] if item['id']==created['id'])
        self.assertEqual(row['name'],'应聘测试')
        self.assertEqual(row['school_name'],'应聘登记测试大学')
        self.assertTrue(row['candidate_id'])
        self.assertTrue(row['resume'])
        self.assertEqual(self.call('applications/'+str(created['id']),{'phone':'13000130999','name':'未匹配登记','school_name':'另一大学'},method='PUT')[0],200)
        row=next(item for item in self.call('applications')[1] if item['id']==created['id'])
        talent=next(item for item in self.call('candidates')[1] if item['id']==row['candidate_id'])
        self.assertEqual(talent['phone'],'13000130999')
        self.assertEqual(talent['name'],'未匹配登记')
        self.assertEqual(row['resume'],'')

    def test_11_identity_derives_application_profile(self):
        from datetime import datetime
        values={'name':'证件测试','school_name':'测试大学','identity_card':'110101200002290026','birthdate':'1999-01-01','gender':'男','age':99}
        status,created=self.call('applications',values)
        self.assertEqual(status,201)
        row=next(item for item in self.call('applications')[1] if item['id']==created['id'])
        self.assertEqual(row['identity_card'],'110101200002290026')
        self.assertEqual(row['birthdate'],'2000-02')
        self.assertEqual(row['gender'],'女')
        now=datetime.now().date()
        self.assertEqual(row['age'],now.year-2000-int(now.month<2))
        self.assertEqual(self.call('applications/'+str(created['id']),{'identity_card':'110101200101020013','gender':'女'},method='PUT')[0],200)
        row=next(item for item in self.call('applications')[1] if item['id']==created['id'])
        self.assertEqual(row['birthdate'],'2001-01')
        self.assertEqual(row['gender'],'男')
        for identity in ('123','110101200102290012','110101299901010012'):
            self.assertEqual(self.call('applications',dict(values,identity_card=identity))[0],400)
        payload=self.file(['姓名','毕业院校','身份证'],[['证件导入','测试大学','110101200101020013']])
        preview=self.call('applications/import-preview',payload)[1]
        self.assertTrue(preview['token'])
        self.assertEqual(preview['rows'][0]['values']['birthdate'],'2001-01')
        self.assertEqual(preview['rows'][0]['values']['gender'],'男')

    def test_12_identity_checksum_validation(self):
        self.assertEqual(server.identity_details('11010519491231002x')['identity_card'],'11010519491231002X')
        values={'name':'号码校验测试','school_name':'测试大学','identity_card':'110101200101020013'}
        self.assertEqual(self.call('applications',values)[0],201)
        invalid=dict(values,identity_card='110101200101020014')
        status,result=self.call('applications',invalid)
        self.assertEqual(status,400)
        self.assertIn('校验码',result['error'])
        preview=self.call('applications/import-preview',self.file(['姓名','毕业院校','身份证'],[['校验错误','测试大学',invalid['identity_card']]]))[1]
        self.assertIsNone(preview['token'])
        self.assertIn('校验码',preview['errors'][0]['error'])
        with self.assertRaises(ValueError):
            server.identity_details('110101200101020000')

    def test_13_public_full_application_and_internal_fields(self):
        fields=self.call('public/schema',public=True)[1]['fields']
        keys={field['key'] for field in fields}
        expected={field['key'] for field in server.SCHEMAS['applications']['fields'] if not field.get('hidden')} - {'screening','needs_first_interview','status','is_elite'}
        self.assertEqual(keys,expected)
        pdf={'name':'学生附件.pdf','content':base64.b64encode(b'%PDF-1.4\nstudent').decode()}
        png={'name':'签名.png','content':'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1kAAAAASUVORK5CYII='}
        payload={'name':'学生完整登记','phone':'13900139999','school_name':'学生大学','major':'土木工程','education':'本科','identity_card':'110101200101020013','birthdate':'1998-01-01','gender':'女','ethnicity':'汉族','hometown':'南京','enrollment_type':'普通全日制','degree':'学士','education_experiences':[{'education':'本科','school_name':'学生大学','major':'土木工程','enrollment':'2022-09','graduation':'2026-07'}],'language':'英语','english_level':'CET-4','other_language':'无','major_category':'工程类','other_major':'工程管理','political_status':'共青团员','health_condition':'否','language_scores':'雅思 6.0','date':'2026-09-07','channel':'校园线下','is_985':'否','is_211':'是','is_double_first':'是','screening':'合格','needs_first_interview':'是','status':'已结束','is_elite':'是','notes':'不得写入','candidate_id':1,'created_date':'伪造','_resume_upload':pdf,'_transcript_upload':pdf,'_certificates_upload':pdf,'_signature_upload':png}
        self.call('candidates',{'name':'预收集','phone':payload['phone'],'school_name':'原学校','major':'原专业','education':'本科'})
        status,result=self.call('public/apply',payload,public=True)
        self.assertEqual(status,201,result)
        application=next(row for row in self.call('applications')[1] if row['id']==result['application_id'])
        candidate=next(row for row in self.call('candidates')[1] if row['id']==result['candidate_id'])
        for key in ('screening','needs_first_interview','is_elite','notes'):
            self.assertEqual(application[key],'')
        self.assertEqual(application['status'],'已生效')
        self.assertEqual(candidate['status'],'待筛选')
        self.assertEqual(application['candidate_id'],candidate['id'])
        self.assertEqual(application['birthdate'],'2001-01')
        self.assertEqual(candidate['birthdate'],'2001-01')
        self.assertEqual(application['cohort'],'2026届')
        for key in ('ethnicity','enrollment_type','degree','language','english_level','other_major','health_condition','is_985','is_211','is_double_first'):
            self.assertEqual(application[key],payload[key])
        for key in ('resume','transcript','certificates','signature'):
            self.assertTrue(application[key])
            self.assertEqual(self.call('attachments/'+application[key],public=True)[0],401)
        self.assertEqual(candidate['resume'],application['resume'])
        self.assertEqual(self.call('public/apply',payload,public=True)[0],200)

    def test_14_application_two_statuses(self):
        payload={'name':'状态测试','school_name':'测试大学'}
        status,created=self.call('applications',payload)
        self.assertEqual(status,201)
        row=next(row for row in self.call('applications')[1] if row['id']==created['id'])
        self.assertEqual(row['status'],'草稿')
        self.assertEqual(self.call('applications/'+str(created['id']),{'status':'已生效'},method='PUT')[0],200)
        self.assertEqual(self.call('applications/'+str(created['id']),{'status':'草稿'},method='PUT')[0],200)
        self.assertEqual(self.call('applications/'+str(created['id']),{'status':'待处理'},method='PUT')[0],400)
        self.assertEqual(self.call('applications/'+str(created['id']),{'status':''},method='PUT')[0],400)

    def test_15_interview_full_profile(self):
        from datetime import datetime
        status,prefill=self.call('interview-prefill?identity_card=110101200101020013')
        self.assertEqual(status,200)
        self.assertTrue(prefill['matched'])
        self.assertEqual(prefill['values']['name'],'学生完整登记')
        values={'identity_card':'110101200101020013','first_interview_consent':'同意','status':'进行中','interview_at':'2026-09-09T14:25:59','is_elite':'是','personality':'沉稳','family':'家庭情况说明','campus_position':'班长','has_internship':'是','project_locations':'华东','other_project_locations':'南京','first_interview_passed':'是','first_interview_notes':'专业基础扎实','second_interview_notes':'沟通良好','result':'发offer','personal_intention':'愿意加入','basic_situation':'应届毕业生','created_date':'伪造'}
        status,created=self.call('interviews',values)
        self.assertEqual(status,201,created)
        row=next(item for item in self.call('interviews')[1] if item['id']==created['id'])
        self.assertEqual(row['interview_at'],'2026-09-09T14:25')
        self.assertEqual(row['name'],'学生完整登记')
        self.assertEqual(row['gender'],'男')
        self.assertEqual(row['birthdate'],'2001-01')
        self.assertNotEqual(row['created_date'],'伪造')
        for key in ('resume','transcript','certificates'):
            self.assertTrue(row[key])
            self.assertEqual(self.call('attachments/'+row[key])[0],200)
        for state in ('草稿','已生效','进行中','已取消'):
            self.assertEqual(self.call('interviews/'+str(row['id']),{'status':state},method='PUT')[0],200)
        self.assertEqual(self.call('interviews/'+str(row['id']),{'result':'pass'},method='PUT')[0],200)
        for updates in ({'result':'通过'},{'first_interview_consent':'是'},{'status':'待面试'},{'interview_at':'2026-02-30T10:00'}):
            self.assertEqual(self.call('interviews/'+str(row['id']),updates,method='PUT')[0],400)
        book=openpyxl.load_workbook(io.BytesIO(self.call('interviews/export.xlsx')[1]));sheet=book.active
        headers=[cell.value for cell in sheet[1]]
        cells=next(cells for cells in sheet.iter_rows(min_row=2) if cells[0].value==row['id'])
        self.assertEqual(cells[headers.index('面试时间')].number_format,'yyyy-mm-dd hh:mm')
        self.assertEqual(cells[headers.index('面试时间')].value,datetime(2026,9,9,14,25))
        self.assertIn('二面情况说明',headers)
        book.close()
        self.assertEqual(self.call('interviews/batch-delete',{'ids':[row['id']]})[0],200)
        self.assertEqual(self.call('attachments/'+row['transcript'])[0],200)
        payload=self.file(['姓名','毕业院校','面试时间','面试结果'],[['独立面试','大学',datetime(2026,9,10,9,40),'pass']])
        preview=self.call('interviews/import-preview',payload)[1]
        self.assertTrue(preview['token'],preview)
        self.assertEqual(self.call('interviews/import-commit',{'token':preview['token']})[0],201)

    def test_16_interview_identity_uses_application_only(self):
        identity='11010519491231002X'
        payload={'name':'独立应聘登记','school_name':'学校甲','phone':'13800138000','identity_card':identity}
        first=self.call('applications',payload)[1]['id']
        latest=self.call('applications',dict(payload,name='最新应聘资料',school_name='学校乙',phone=''))[1]['id']
        self.assertGreater(latest,first)
        status,result=self.call('interview-prefill?identity_card='+identity.lower())
        self.assertEqual(status,200)
        self.assertEqual(result['source_id'],latest)
        self.assertEqual(result['values']['name'],'最新应聘资料')
        status,detail=self.call('student-registration?identity_card='+identity.lower())
        self.assertEqual(status,200)
        self.assertEqual(detail['values']['name'],'最新应聘资料')
        self.assertEqual(detail['values']['school_name'],'学校乙')
        self.assertNotIn('status',detail['values'])
        self.assertNotIn('screening',detail['values'])
        self.assertNotIn('candidate_id',detail['values'])
        self.assertEqual(self.call('student-registration?identity_card='+identity,public=True)[0],401)
        self.assertFalse(self.call('student-registration?phone=13800138000')[1]['matched'])
        self.assertEqual(self.call('student-registration?identity_card=invalid')[0],400)
        self.assertFalse(self.call('interview-prefill?phone=13800138000')[1]['matched'])
        self.assertEqual(self.call('interview-prefill?identity_card='+identity,public=True)[0],401)
        status,created=self.call('interviews',{'identity_card':identity})
        self.assertEqual(status,201)
        row=next(row for row in self.call('interviews')[1] if row['id']==created['id'])
        self.assertEqual(row['name'],'最新应聘资料')
        self.assertEqual(row['school_name'],'学校乙')
        self.assertIsNone(row['candidate_id'])
        self.assertEqual(row['phone'],'')
        self.assertEqual(self.call('interviews',{'phone':'13800138000'})[0],400)

    def test_17_employee_profile(self):
        payload={'name':'新员工测试','identity_card':'11010519491231002x','gender':'男','birthdate':'2000-01-01','ethnicity':'汉族','phone':'13912345678','school_name':'测试大学','education':'本科','major':'工程','is_elite':'是','height':'178.5','weight':'70','shoe_size':'42.5','position':'技术员','home_address':'测试地址','enrollment_type':'普通全日制','language':'英语','channel':'校园线下','hometown':'南京','notes':'备注测试'}
        status,result=self.call('employees',payload)
        self.assertEqual(status,201)
        row=next(row for row in self.call('employees')[1] if row['id']==result['id'])
        self.assertEqual(row['identity_card'],'11010519491231002X')
        self.assertEqual(row['birthdate'],'1949-12')
        self.assertEqual(row['gender'],'女')
        for key in ('height','weight','shoe_size','home_address','channel','notes'):
            self.assertEqual(row[key],payload[key])
        self.assertEqual(self.call('employees',dict(payload,identity_card='invalid'))[0],400)
        self.assertEqual(self.call('employees',dict(payload,height='-1'))[0],400)
        self.assertEqual(self.call('employees',dict(payload,channel='其他'))[0],400)

    def test_18_employee_identity_prefill(self):
        identity='11010519491231002X'
        self.call('applications',{'name':'新员工来源','school_name':'来源大学','identity_card':identity,'ethnicity':'汉族','enrollment_type':'统招专升本','channel':'智联招聘','is_elite':'是'})
        status,result=self.call('employee-prefill?identity_card='+identity.lower())
        self.assertEqual(status,200)
        self.assertEqual(result['values']['name'],'新员工来源')
        self.assertEqual(result['values']['enrollment_type'],'统招专升本')
        self.assertEqual(result['values']['channel'],'智联招聘')
        self.assertNotIn('candidate_id',result['values'])
        self.assertNotIn('notes',result['values'])
        self.assertEqual(self.call('employee-prefill?identity_card='+identity,public=True)[0],401)
        self.assertFalse(self.call('employee-prefill?phone=13800138000')[1]['matched'])
        self.assertEqual(self.call('employee-prefill?identity_card=invalid')[0],400)
        status,created=self.call('employees',{'identity_card':identity,'weight':'65'})
        self.assertEqual(status,201)
        row=next(row for row in self.call('employees')[1] if row['id']==created['id'])
        self.assertEqual(row['name'],'新员工来源')
        self.assertEqual(row['ethnicity'],'汉族')
        self.assertEqual(row['weight'],'65')

    def test_19_public_employee_registration(self):
        status,schema=self.call('public/employee-schema',public=True)
        self.assertEqual(status,200)
        self.assertEqual(len(schema['fields']),22)
        self.assertNotIn('candidate_id',[field['key'] for field in schema['fields']])
        payload={'name':'外链新员工','identity_card':'110101200101020013','phone':'13987654321','height':'180','channel':'校园平台','candidate_id':99999,'status':'已报到'}
        status,result=self.call('public/employees',payload,public=True)
        self.assertEqual(status,201)
        self.assertEqual(result,{'ok':True})
        row=next(row for row in self.call('employees')[1] if row['name']=='外链新员工')
        self.assertIsNone(row['candidate_id'])
        self.assertEqual(row['status'],'')
        self.assertEqual(row['birthdate'],'2001-01')
        self.assertEqual(row['height'],'180')
        self.assertEqual(self.call('public/employees',payload,public=True)[0],409)
        self.assertEqual(self.call('public/employees',dict(payload,identity_card='invalid'),public=True)[0],400)
        self.assertEqual(self.call('public/employees',dict(payload,phone=''),public=True)[0],400)
        self.assertEqual(self.call('employees',public=True)[0],401)

    def test_20_language_summary_prefill(self):
        identity='11010519491231002X'
        self.call('applications',{'name':'语言汇总测试','school_name':'测试大学','identity_card':identity,'language':'英语','english_level':'CET-4','other_language':'日语N2','language_scores':'雅思7分'})
        expected='语言能力：英语；英语等级：CET-4；其他语言能力：日语N2；雅思/托福成绩：雅思7分'
        for endpoint,entity in (('interview-prefill','interviews'),('employee-prefill','employees')):
            self.assertEqual(self.call(endpoint+'?identity_card='+identity)[1]['values']['language'],expected)
            status,created=self.call(entity,{'identity_card':identity})
            self.assertEqual(status,201)
            row=next(row for row in self.call(entity)[1] if row['id']==created['id'])
            self.assertEqual(row['language'],expected)
        self.assertEqual(server.language_summary({'language':'无','english_level':'','other_language':None,'language_scores':''}),'语言能力：无')
        self.assertEqual(server.language_summary({'language':'','english_level':'','other_language':'','language_scores':''}),'')

    def test_22_cascade_deletion(self):
        identity='110101200101020013'
        candidate=self.call('candidates',{'name':'级联测试','phone':'13912349876','school_name':'测试大学','major':'工程','education':'本科'})[1]['id']
        application=self.call('applications',{'name':'级联测试','phone':'13912349876','school_name':'测试大学','identity_card':identity})[1]['id']
        interview=self.call('interviews',{'identity_card':identity})[1]['id']
        employee=self.call('employees',{'identity_card':identity})[1]['id']
        preview=self.call('candidates/delete-preview',{'ids':[candidate]})[1]
        keys={(row['entity'],row['id']) for row in preview['items']}
        self.assertTrue({('candidates',candidate),('applications',application),('interviews',interview),('employees',employee)} <= keys)
        self.assertEqual(self.call('candidates/batch-delete',{'ids':[candidate],'expected':[]})[0],400)
        self.assertTrue(any(row['id']==candidate for row in self.call('candidates')[1]))
        result=self.call('candidates/batch-delete',{'ids':[candidate],'expected':preview['items']})
        self.assertEqual(result[0],200)
        self.assertEqual(result[1]['count'],preview['count'])
        for entity,identifier in keys:
            self.assertFalse(any(row['id']==identifier for row in self.call(entity)[1]))

    def test_33_candidate_education_subtable(self):
        payload={'name':'多学历测试','phone':'13966665555','education_experiences':[
            {'education':'本科','school_name':'本科大学','major':'土木工程'},
            {'education':'硕士','school_name':'硕士大学','major':'结构工程'}]}
        status,created=self.call('candidates',payload)
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['education'],'硕士')
        self.assertEqual(row['school_name'],'硕士大学')
        self.assertEqual(row['major'],'结构工程')
        self.assertEqual(len(row['education_experiences']),2)
        prefill=self.call('application-prefill?phone=13966665555')[1]
        self.assertEqual((prefill['values']['education'],prefill['values']['school_name'],prefill['values']['major']),('硕士','硕士大学','结构工程'))
        updated={'education_experiences':[{'education':'博士','school_name':'博士大学','major':'工程管理','enrollment':'','graduation':''}]}
        self.assertEqual(self.call('candidates/'+str(created['id']),updated,method='PUT')[0],200)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['education_experiences'],updated['education_experiences'])
        self.assertEqual((row['education'],row['school_name'],row['major']),('博士','博士大学','工程管理'))
        self.assertEqual(self.call('candidates',{'name':'错误学历','phone':'13966665556','education_experiences':[{'education':'硕士','school_name':'','major':'工程'}]})[0],400)

    def test_34_education_subtables_in_other_records(self):
        histories=[
            {'education':'本科','school_name':'本科院校','major':'机械工程','enrollment':'2018-09','graduation':'2022-06'},
            {'education':'博士','school_name':'博士院校','major':'化学工程','enrollment':'2024-09','graduation':'2028-06'},
        ]
        identity='110101200101020013'
        status,application=self.call('applications',{'name':'多学历联动','phone':'13966665557','identity_card':identity,'education_experiences':histories})
        self.assertEqual(status,201)
        application_row=next(row for row in self.call('applications')[1] if row['id']==application['id'])
        self.assertEqual(application_row['education_experiences'],histories)
        self.assertEqual((application_row['education'],application_row['school_name'],application_row['major']),('博士','博士院校','化学工程'))
        candidate=next(row for row in self.call('candidates')[1] if row['id']==application_row['candidate_id'])
        self.assertEqual(candidate['education_experiences'],histories)
        interview_prefill=self.call('interview-prefill?identity_card='+identity)[1]
        employee_prefill=self.call('employee-prefill?identity_card='+identity)[1]
        self.assertEqual(interview_prefill['values']['education_experiences'],histories)
        self.assertEqual(employee_prefill['values']['education_experiences'],histories)
        status,interview=self.call('interviews',{'name':'多学历联动','identity_card':identity,'education_experiences':histories})
        self.assertEqual(status,201)
        interview_row=next(row for row in self.call('interviews')[1] if row['id']==interview['id'])
        self.assertEqual((interview_row['education'],interview_row['school_name'],interview_row['major']),('博士','博士院校','化学工程'))
        status,employee=self.call('employees',{'name':'多学历联动','identity_card':identity,'phone':'13966665557','education_experiences':histories})
        self.assertEqual(status,201)
        employee_row=next(row for row in self.call('employees')[1] if row['id']==employee['id'])
        self.assertEqual(employee_row['education_experiences'],histories)
        self.assertEqual((employee_row['education'],employee_row['school_name'],employee_row['major']),('博士','博士院校','化学工程'))
        public_histories=[
            {'education':'本科','school_name':'公开本科院校','major':'过程装备','enrollment':'2019-09','graduation':'2023-06'},
            {'education':'硕士','school_name':'公开硕士院校','major':'化工机械','enrollment':'2023-09','graduation':'2026-06'},
        ]
        status,result=self.call('public/apply',{'name':'外链多学历','phone':'13966665558','education_experiences':public_histories},public=True)
        self.assertEqual(status,201)
        public_application=next(row for row in self.call('applications')[1] if row['id']==result['application_id'])
        self.assertEqual(public_application['education_experiences'],public_histories)
        public_candidate=next(row for row in self.call('candidates')[1] if row['id']==result['candidate_id'])
        self.assertEqual(public_candidate['education_experiences'],public_histories)
        status,_=self.call('public/employees',{'name':'外链新员工多学历','identity_card':'320101200201010014','phone':'13966665559','education_experiences':public_histories},public=True)
        self.assertEqual(status,201)
        public_employee=next(row for row in self.call('employees')[1] if row['identity_card']=='320101200201010014')
        self.assertEqual(public_employee['education_experiences'],public_histories)

    def test_35_education_dates_drive_candidate_graduation_and_campus_cohort(self):
        histories=[
            {'education':'本科','school_name':'本科大学','major':'土木工程','enrollment':'2018-09','graduation':'2022-06'},
            {'education':'硕士','school_name':'硕士大学','major':'结构工程','enrollment':'2022-09','graduation':'2025-06'},
        ]
        status,created=self.call('candidates',{'name':'校园招聘应届测试','phone':'13966665560','recruitment_type':'校园招聘','education_experiences':histories})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['education_experiences'],histories)
        self.assertEqual(row['graduation'],'2025-06')
        self.assertEqual(row['cohort'],'2025届')

        status,created=self.call('candidates',{'name':'社会招聘届别测试','phone':'13966665561','recruitment_type':'社会招聘','education_experiences':histories})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['education_experiences'],histories)
        self.assertEqual(row['graduation'],'')
        self.assertEqual(row['cohort'],'')

        invalid=[{'education':'本科','school_name':'日期错误大学','major':'工程','enrollment':'2025-09','graduation':'2025-06'}]
        self.assertEqual(self.call('candidates',{'name':'日期错误','phone':'13966665562','recruitment_type':'校园招聘','education_experiences':invalid})[0],400)

    def test_36_parent_time_fields_are_removed_from_non_candidate_schemas(self):
        schemas=self.call('schema')[1]
        candidate_fields={field['key']:field for field in schemas['candidates']['fields']}
        self.assertTrue(candidate_fields['graduation']['readonly'])
        self.assertNotIn('experience',candidate_fields)
        self.assertNotIn('enrollment',{field['key'] for field in schemas['applications']['fields']})
        self.assertNotIn('graduation',{field['key'] for field in schemas['applications']['fields']})
        self.assertNotIn('graduation',{field['key'] for field in schemas['interviews']['fields']})

    def test_37_explicit_empty_education_history_clears_optional_record(self):
        history=[{'education':'本科','school_name':'清空前大学','major':'工程','enrollment':'2018-09','graduation':'2022-06'}]
        status,created=self.call('employees',{'name':'清空教育经历','education_experiences':history})
        self.assertEqual(status,201)
        status,_=self.call('employees/'+str(created['id']),{'name':'清空教育经历','education_experiences':[]},method='PUT')
        self.assertEqual(status,200)
        row=next(row for row in self.call('employees')[1] if row['id']==created['id'])
        self.assertEqual(row['education_experiences'],[])
        self.assertEqual((row['education'],row['school_name'],row['major']),('','',''))

    def test_38_every_required_candidate_education_row_is_validated(self):
        histories=[
            {'education':'本科','school_name':'完整大学','major':'工程','enrollment':'2018-09','graduation':'2022-06'},
            {'education':'硕士','school_name':'','major':'结构工程','enrollment':'2022-09','graduation':'2025-06'},
        ]
        status,result=self.call('candidates',{'name':'第二学历不完整','phone':'13966665564','education_experiences':histories})
        self.assertEqual(status,400)
        self.assertIn('第2条教育经历',result['error'])

    def test_39_work_experience_is_shared_by_both_recruitment_types(self):
        education=[{'education':'本科','school_name':'经历测试大学','major':'工程','enrollment':'2018-09','graduation':'2022-06'}]
        work=[{'organization':'社会招聘单位','position':'项目工程师','start_date':'2022-07','end_date':'2025-08','description':'负责现场管理'}]
        status,created=self.call('candidates',{'name':'经历通用测试','phone':'13966665565','recruitment_type':'社会招聘','education_experiences':education,'work_experiences':work})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['work_experiences'],work)
        self.assertNotIn('internship_experiences',row)
        self.assertNotIn('experience',row)

        status,_=self.call('candidates/'+str(created['id']),{'recruitment_type':'校园招聘'},method='PUT')
        self.assertEqual(status,200)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['work_experiences'],work)

        campus_work=[{'organization':'校园招聘工作单位','position':'实践岗位','start_date':'2021-07','end_date':'','description':'校园招聘人才工作经历'}]
        status,campus=self.call('candidates',{'name':'校园招聘工作经历','phone':'13966665569','recruitment_type':'校园招聘','education_experiences':education,'work_experiences':campus_work})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==campus['id'])
        self.assertEqual(row['work_experiences'],campus_work)
        self.assertNotIn('internship_experiences',row)

    def test_40_candidate_experience_dates_must_be_valid_and_ordered(self):
        education=[{'education':'本科','school_name':'经历日期大学','major':'工程','enrollment':'2018-09','graduation':'2022-06'}]
        invalid=[{'organization':'时间错误单位','position':'工程师','start_date':'2025-09','end_date':'2025-06','description':''}]
        status,result=self.call('candidates',{'name':'经历日期错误','phone':'13966665566','recruitment_type':'社会招聘','education_experiences':education,'work_experiences':invalid})
        self.assertEqual(status,400)
        self.assertIn('开始时间不能晚于结束时间',result['error'])

        optional_position=[{'organization':'岗位选填单位','position':'','start_date':'2025-01','end_date':'','description':''}]
        status,created=self.call('candidates',{'name':'工作岗位选填','phone':'13966665567','recruitment_type':'社会招聘','education_experiences':education,'work_experiences':optional_position})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['work_experiences'],optional_position)

        missing_organization=[{'organization':'','position':'工程师','start_date':'2025-01','end_date':'','description':''}]
        status,result=self.call('candidates',{'name':'工作单位仍必填','phone':'13966665575','recruitment_type':'社会招聘','education_experiences':education,'work_experiences':missing_organization})
        self.assertEqual(status,400)
        self.assertIn('单位',result['error'])

        malformed=[{'organization':{'name':'对象单位'},'position':'工程师','start_date':'','end_date':'','description':''}]
        status,result=self.call('candidates',{'name':'经历格式错误','phone':'13966665568','recruitment_type':'社会招聘','education_experiences':education,'work_experiences':malformed})
        self.assertEqual(status,400)
        self.assertIn('格式不正确',result['error'])

    def test_41_project_experience_is_social_only_and_deleted_on_type_change(self):
        education=[{'education':'本科','school_name':'项目经历大学','major':'工程','enrollment':'2018-09','graduation':'2022-06'}]
        projects=[{'project_name':'化工建设项目','role':'项目负责人','start_date':'2022-07','end_date':'2025-08','description':'负责项目全过程管理'}]
        status,created=self.call('candidates',{'name':'社会招聘项目经历','phone':'13966665570','recruitment_type':'社会招聘','education_experiences':education,'project_experiences':projects})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['project_experiences'],projects)

        status,_=self.call('candidates/'+str(created['id']),{'recruitment_type':'校园招聘'},method='PUT')
        self.assertEqual(status,200)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertNotIn('project_experiences',row)

        status,campus=self.call('candidates',{'name':'校园招聘无项目经历','phone':'13966665571','recruitment_type':'校园招聘','education_experiences':education,'project_experiences':projects})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==campus['id'])
        self.assertNotIn('project_experiences',row)

    def test_42_project_experience_validates_required_fields_and_dates(self):
        education=[{'education':'本科','school_name':'项目校验大学','major':'工程','enrollment':'2018-09','graduation':'2022-06'}]
        invalid=[{'project_name':'日期错误项目','role':'负责人','start_date':'2025-09','end_date':'2025-06','description':''}]
        status,result=self.call('candidates',{'name':'项目日期错误','phone':'13966665572','recruitment_type':'社会招聘','education_experiences':education,'project_experiences':invalid})
        self.assertEqual(status,400)
        self.assertIn('开始时间不能晚于结束时间',result['error'])

        optional_role=[{'project_name':'角色选填项目','role':'','start_date':'2025-01','end_date':'','description':''}]
        status,created=self.call('candidates',{'name':'项目角色选填','phone':'13966665573','recruitment_type':'社会招聘','education_experiences':education,'project_experiences':optional_role})
        self.assertEqual(status,201)
        row=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(row['project_experiences'],optional_role)

        missing_project_name=[{'project_name':'','role':'负责人','start_date':'2025-01','end_date':'','description':''}]
        status,result=self.call('candidates',{'name':'项目名称仍必填','phone':'13966665576','recruitment_type':'社会招聘','education_experiences':education,'project_experiences':missing_project_name})
        self.assertEqual(status,400)
        self.assertIn('项目名称',result['error'])

        for suffix,start_date in [('无开始时间',''),('非法日期','2025-02-99')]:
            invalid=[{'project_name':suffix+'项目','role':'负责人','start_date':start_date,'end_date':'','description':''}]
            status,result=self.call('candidates',{'name':suffix,'phone':'13966665574','recruitment_type':'社会招聘','education_experiences':education,'project_experiences':invalid})
            self.assertEqual(status,400)
            self.assertIn('开始时间',result['error'])
        self.assertFalse(any(row['phone']=='13966665574' for row in self.call('candidates')[1]))

    def test_43_project_experience_update_import_export_and_cascade_contracts(self):
        education=[{'education':'本科','school_name':'项目契约大学','major':'工程','enrollment':'2018-09','graduation':'2022-06'}]
        projects=[{'project_name':'保留项目','role':'负责人','start_date':'2022-07','end_date':'','description':'持续负责'}]
        status,created=self.call('candidates',{'name':'项目契约人才','phone':'13966665575','recruitment_type':'社会招聘','education_experiences':education,'project_experiences':projects})
        self.assertEqual(status,201)
        identifier=created['id']
        self.assertEqual(self.call('candidates/'+str(identifier),{'name':'项目契约人才（更新）'},method='PUT')[0],200)
        row=next(row for row in self.call('candidates')[1] if row['id']==identifier)
        self.assertEqual(row['project_experiences'],projects)
        self.assertEqual(self.call('candidates/'+str(identifier),{'project_experiences':[]},method='PUT')[0],200)
        row=next(row for row in self.call('candidates')[1] if row['id']==identifier)
        self.assertEqual(row['project_experiences'],[])

        self.assertEqual(self.call('candidates/'+str(identifier),{'project_experiences':projects},method='PUT')[0],200)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM candidate_project_experiences WHERE candidate_id=?',(identifier,)).fetchone()[0],1)
        preview=self.call('candidates/delete-preview',{'ids':[identifier]})[1]
        self.assertEqual(self.call('candidates/batch-delete',{'ids':[identifier],'expected':preview['items']})[0],200)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM candidate_project_experiences WHERE candidate_id=?',(identifier,)).fetchone()[0],0)

        pdf={'name':'项目经历导入.pdf','content':base64.b64encode(b'%PDF-1.4\nproject-import').decode()}
        document=self.call('resume_documents',{'name':'项目经历导入','_resume_upload':pdf})[1]['id']
        self.assertEqual(self.call('resume_documents/'+str(document),{'recruitment_type':'社会招聘'},method='PUT')[0],200)
        status,created=self.call('resume_documents/'+str(document)+'/candidate',{'name':'项目导入人才','phone':'13966665576','education_experiences':education,'project_experiences':projects})
        self.assertEqual(status,201)
        imported=next(row for row in self.call('candidates')[1] if row['id']==created['id'])
        self.assertEqual(imported['project_experiences'],projects)

        book=openpyxl.load_workbook(io.BytesIO(self.call('candidates/export.xlsx')[1]))
        self.assertNotIn('项目经历',[cell.value for cell in book.active[1]])
        book.close()

    def test_32_resume_recognition(self):
        payload={'recruitment_type':'校园招聘','channel':'校园平台','_resume_upload':{'name':'批量识别测试.pdf','content':base64.b64encode(b'%PDF-1.4\nrecognition-test').decode()}}
        status,created=self.call('resume_documents',payload)
        self.assertEqual(status,201)
        path='resume_documents/'+str(created['id'])
        self.assertEqual(self.call(path+'/recognize',{},public=True)[0],401)
        with patch('server.dify_client.configuration',side_effect=ValueError('缺少配置')):
            self.assertFalse(self.call('recognition-config')[1]['configured'])
            self.assertEqual(self.call(path+'/recognize',{})[0],400)
        with patch('server.dify_client.configuration',return_value={}),patch('server.dify_client.recognize',side_effect=ValueError('识别失败')):
            self.assertEqual(self.call(path+'/recognize',{})[0],400)
        row=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        self.assertEqual(row['status'],'失败')
        before=len(self.call('candidates')[1])
        workflow_result={'recruitment_type':'校园招聘','position_name':'施工员','candidate':{
            'basic':{'name':'识别测试','phone':'13977776666'},
            'education':[{'school':'识别大学','major':'土木工程','degree':'本科','start_date':'2022-09','end_date':'2026-06'}],
            'work_experience':[{'company':'实习单位','position':'实习生','start_date':'2025-01','end_date':'至今'}],
            'projects':[{'project_name':'校园项目','role':'成员','start_date':'2025-01','end_date':'至今'}]
        }}
        with patch('server.dify_client.configuration',return_value={}),patch('server.dify_client.recognize',return_value=workflow_result) as run:
            self.assertEqual(self.call(path+'/recognize',{})[0],200)
            self.assertEqual(self.call(path+'/recognize',{})[0],409)
            self.assertEqual(run.call_count,1)
        self.assertEqual(run.call_args.args[2]['recruitment_type'],'校园招聘')
        recognized=self.call(path+'/recognition-result')[1]['result']
        self.assertEqual(recognized['name'],'识别测试')
        self.assertEqual(recognized['education_experiences'][0]['school_name'],'识别大学')
        self.assertEqual(recognized['work_experiences'][0]['end_date'],'')
        self.assertNotIn('project_experiences',recognized)
        self.assertEqual(self.call(path+'/candidate',{'name':'缺少必填字段'})[0],400)
        self.assertEqual(next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])['status'],'已识别')
        self.assertEqual(self.call(path+'/recognition-result')[1]['result']['name'],'识别测试')
        status,created_candidate=self.call(path+'/candidate',recognized)
        self.assertEqual(status,201)
        candidate=next(row for row in self.call('candidates')[1] if row['id']==created_candidate['id'])
        self.assertEqual(candidate['education_experiences'][0]['school_name'],'识别大学')
        self.assertEqual(candidate['work_experiences'][0]['organization'],'实习单位')
        self.assertNotIn('project_experiences',candidate)
        self.assertEqual(len(self.call('candidates')[1]),before+1)
        self.assertEqual(next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])['status'],'已入库')

        race_payload={'recruitment_type':'校园招聘','channel':'校园平台','_resume_upload':{'name':'并发修改识别.pdf','content':base64.b64encode(b'%PDF-1.4\nrecognition-race').decode()}}
        race_id=self.call('resume_documents',race_payload)[1]['id']
        race_path='resume_documents/'+str(race_id)
        def edit_during_recognition(*_):
            with server.db() as conn:
                conn.execute("UPDATE resume_documents SET channel='校园线下' WHERE id=?",(race_id,))
            return workflow_result
        with patch('server.dify_client.configuration',return_value={}),patch('server.dify_client.recognize',side_effect=edit_during_recognition):
            self.assertEqual(self.call(race_path+'/recognize',{})[0],409)
        race_row=next(row for row in self.call('resume_documents')[1] if row['id']==race_id)
        self.assertEqual(race_row['channel'],'校园线下')
        self.assertEqual(race_row['status'],'待识别')
        self.assertIsNone(self.call(race_path+'/recognition-result')[1]['result'])

    def test_31_overview_recruitment_groups(self):
        for kind,phone in [('校园招聘','13988886661'),('社会招聘','13988886662')]:
            self.assertEqual(self.call('candidates',{'name':kind+'统计测试','phone':phone,'school_name':'测试大学','major':'工程','education':'本科','recruitment_type':kind})[0],201)
        stats=self.call('stats')[1]
        with server.db() as conn:
            for kind in ('校园招聘','社会招聘'):
                group=stats['recruitment'][kind]
                for key in ('candidates','resume_documents'):
                    expected=conn.execute(f'SELECT COUNT(*) FROM {key} WHERE recruitment_type=?',(kind,)).fetchone()[0]
                    self.assertEqual(group['counts'][key],expected)
                self.assertNotIn('employees',group['counts'])
                for row in group['recent']:
                    self.assertEqual(conn.execute('SELECT recruitment_type FROM candidates WHERE id=?',(row['id'],)).fetchone()[0],kind)
        for key in ('applications','interviews','schools'):
            self.assertEqual(stats['recruitment']['校园招聘']['counts'][key],stats['counts'][key])
            self.assertNotIn(key,stats['recruitment']['社会招聘']['counts'])

    def test_30_application_upserts_talent(self):
        payload={'name':'更新测试','phone':'13988887777','school_name':'测试大学','major':'工程','education':'本科'}
        before=len(self.call('candidates')[1])
        status,unmatched=self.call('public/apply',payload,public=True)
        self.assertEqual(status,201)
        candidate=unmatched['candidate_id']
        self.assertIsInstance(candidate,int)
        self.assertEqual(len(self.call('candidates')[1]),before+1)
        prefill=self.call('public/application-prefill?phone='+payload['phone'],public=True)[1]
        self.assertEqual(prefill['values']['name'],'更新测试')
        self.assertNotIn('resume',prefill['values'])
        self.assertNotIn('status',prefill['values'])
        status,result=self.call('public/apply',dict(payload,name='学生修改',major='自动化'),public=True)
        self.assertEqual(status,200)
        self.assertEqual(result['candidate_id'],candidate)
        self.assertEqual(result['application_id'],unmatched['application_id'])
        row=next(row for row in self.call('candidates')[1] if row['id']==candidate)
        self.assertEqual((row['name'],row['major']),('学生修改','自动化'))
        self.assertEqual(len(self.call('candidates')[1]),before+1)
        self.assertEqual(self.call('applications/'+str(result['application_id']),{'name':'成员修改'},method='PUT')[0],200)
        self.assertEqual(next(row for row in self.call('candidates')[1] if row['id']==candidate)['name'],'成员修改')
        self.assertEqual(len(self.call('candidates')[1]),before+1)
        status,created=self.call('applications',dict(payload,phone='13988887778',name='管理端新增'))
        self.assertEqual(status,201)
        application=next(row for row in self.call('applications')[1] if row['id']==created['id'])
        talent=next(row for row in self.call('candidates')[1] if row['id']==application['candidate_id'])
        self.assertEqual(talent['name'],'管理端新增')
        self.assertEqual(talent['archive_no'],application['archive_no'])
        self.assertEqual(len(self.call('candidates')[1]),before+2)

    def test_29_recruitment_channels(self):
        for entity in ('candidates','resume_documents'):
            with server.db() as conn:
                base={'name':'渠道测试','phone':'13900002222','school_name':'测试大学','major':'工程','education':'本科'}
                for kind,channels in [('校园招聘',['校园线下','校园平台','智联招聘']),('社会招聘',['boss直聘','智联招聘','化工英才网','其他'])]:
                    for channel in channels:
                        result=server.validate(entity,dict(base,recruitment_type=kind,channel=channel,channel_detail='渠道说明'),conn)
                        self.assertEqual(result['channel'],channel)
                        self.assertEqual(result['channel_detail'],'渠道说明' if channel=='其他' else '')
                with self.assertRaises(ValueError):server.validate(entity,dict(base,recruitment_type='校园招聘',channel='boss直聘'),conn)
                with self.assertRaises(ValueError):server.validate(entity,dict(base,recruitment_type='社会招聘',channel='校园线下'),conn)
                with self.assertRaises(ValueError):server.validate(entity,dict(base,recruitment_type='社会招聘',channel='其他'),conn)
        payload={'channel':'其他','channel_detail':'行业交流群','_resume_upload':{'name':'社会招聘渠道.pdf','content':base64.b64encode(b'%PDF-1.4\nsocial-channel-test').decode()}}
        self.assertEqual(self.call('public/resumes/social',payload,public=True)[0],201)
        source=next(row for row in self.call('resume_documents')[1] if row['name']=='社会招聘渠道.pdf')
        self.assertEqual((source['channel'],source['channel_detail']),('其他','行业交流群'))
        self.assertEqual(self.call('resume_documents/'+str(source['id']),{'channel':'其他','channel_detail':'行业交流群'},method='PUT')[0],200)
        result=self.call('resume_documents/'+str(source['id'])+'/candidate',{'name':'社会招聘人才','phone':'13900002222','school_name':'测试大学','major':'工程','education':'本科'})
        self.assertEqual(result[0],201)
        candidate=next(row for row in self.call('candidates')[1] if row['id']==result[1]['id'])
        self.assertEqual((candidate['recruitment_type'],candidate['channel'],candidate['channel_detail']),('社会招聘','其他','行业交流群'))
        invalid={'channel':'校园线下','_resume_upload':{'name':'错误渠道.pdf','content':base64.b64encode(b'%PDF-1.4\ninvalid-channel').decode()}}
        self.assertEqual(self.call('public/resumes/social',invalid,public=True)[0],400)

    def test_28_public_resume_qr(self):
        self.assertEqual(self.call('resume-collection-qr',public=True)[0],401)
        qr=self.call('resume-collection-qr')[1]['items']
        self.assertEqual([item['label'] for item in qr],['校园招聘','社会招聘'])
        self.assertTrue(qr[0]['url'].endswith('/resume-submit?type=campus'))
        self.assertTrue(qr[1]['url'].endswith('/resume-submit?type=social'))
        fields=self.call('public/resume-schema',public=True)[1]['fields']
        self.assertEqual({field['key'] for field in fields},{'position','position_detail','resume','channel','channel_detail'})
        for kind,label in [('campus','校园招聘'),('social','社会招聘')]:
            payload={'name':'扫码'+kind,'position':'手工填写岗位','notes':'测试','recruitment_type':'伪造','status':'已入库','archive_no':'fake','source_id':'fake','_resume_upload':{'name':kind+'.pdf','content':base64.b64encode(b'%PDF-1.4\npublic-'+kind.encode()).decode()}}
            self.assertEqual(self.call('public/resumes/'+kind,payload,public=True)[0],201)
            row=next(row for row in self.call('resume_documents')[1] if row['name']==kind+'.pdf')
            self.assertEqual(row['recruitment_type'],label)
            if kind=='campus':self.assertEqual(row['channel'],'校园线下')
            self.assertEqual(row['status'],'待识别')
            self.assertEqual(row['created_by'],'公开投递')
            self.assertEqual(row['position'],'其它')
            self.assertEqual(row['position_detail'],'手工填写岗位')
            self.assertEqual(row['archive_no'],row['source_id'])
            self.assertNotEqual(row['source_id'],'fake')
            self.assertEqual(self.call('public/resumes/'+kind,payload,public=True)[0],400)
            self.assertEqual(self.call('attachments/'+row['resume'],public=True)[0],401)
        self.assertEqual(self.call('public/resumes/campus',{'name':'缺附件'},public=True)[0],400)

        inferred={'channel':'校园平台','_resume_upload':{'name':'张坤洋_26岁_电气工程师_南京_智联简历.pdf','content':base64.b64encode(b'%PDF-1.4\npublic-inferred-position').decode()}}
        self.assertEqual(self.call('public/resumes/campus',inferred,public=True)[0],201)
        row=next(row for row in self.call('resume_documents')[1] if row['name']==inferred['_resume_upload']['name'])
        self.assertEqual(row['position'],'电气工程师')

        image={'channel':'校园平台','_resume_upload':{'name':'图片简历.PNG','content':base64.b64encode(b'\x89PNG\r\n\x1a\npublic-image-resume').decode()}}
        self.assertEqual(self.call('public/resumes/campus',image,public=True)[0],201)
        row=next(row for row in self.call('resume_documents')[1] if row['name']=='图片简历.PNG')
        self.assertEqual(self.call('attachments/'+row['resume']+'?preview=1')[0],200)

    def test_27_daily_archive_sequences_are_stable_and_not_reused(self):
        with server.db() as conn:
            ids=[]
            for i in range(4):
                identifier=server.insert(conn,'employees',server.validate('employees',{'name':'同秒测试'+str(i)},conn))
                conn.execute("UPDATE employees SET created_at='2030-07-01T18:23:08',archive_no='' WHERE id=?",(identifier,))
                ids.append(identifier)
            server.synchronize_archives(conn)
            numbers=[conn.execute('SELECT archive_no FROM employees WHERE id=?',(identifier,)).fetchone()[0] for identifier in ids]
            self.assertEqual(numbers,['0120300701001','0120300701002','0120300701003','0120300701004'])
            server.synchronize_archives(conn)
            self.assertEqual([conn.execute('SELECT archive_no FROM employees WHERE id=?',(identifier,)).fetchone()[0] for identifier in ids],numbers)
            conn.execute('DELETE FROM employees WHERE id=?',(ids[-1],))
            new_id=server.insert(conn,'employees',server.validate('employees',{'name':'后续同秒'},conn))
            conn.execute("UPDATE employees SET created_at='2030-07-01T18:23:08',archive_no='' WHERE id=?",(new_id,))
            server.synchronize_archives(conn)
            self.assertEqual(conn.execute('SELECT archive_no FROM employees WHERE id=?',(new_id,)).fetchone()[0],'0120300701005')

    def test_26_shared_archive_numbers(self):
        pdf={'name':'编号.pdf','content':base64.b64encode(b'%PDF-1.4\narchive-number-test').decode()}
        document=self.call('resume_documents',{'name':'档案编号测试','_resume_upload':pdf})[1]['id']
        with server.db() as conn:
            conn.execute("UPDATE resume_documents SET created_at='2026-07-01T18:23:08',archive_no='' WHERE id=?",(document,))
            server.synchronize_archives(conn)
        source=next(row for row in self.call('resume_documents')[1] if row['id']==document)
        self.assertEqual(source['archive_no'],'0120260701001')
        candidate=self.call('resume_documents/'+str(document)+'/candidate',{'name':'档案编号测试','phone':'13977778888','school_name':'测试大学','major':'工程','education':'本科','archive_no':'伪造'})[1]['id']
        application=self.call('applications',{'phone':'13977778888','identity_card':'110101200101020013'})[1]['id']
        interview=self.call('interviews',{'identity_card':'110101200101020013'})[1]['id']
        employee=self.call('employees',{'identity_card':'110101200101020013'})[1]['id']
        for entity,identifier in [('candidates',candidate),('applications',application),('interviews',interview),('employees',employee)]:
            row=next(row for row in self.call(entity)[1] if row['id']==identifier)
            self.assertEqual(row['archive_no'],source['archive_no'])
        self.assertEqual(self.call('application-prefill?phone=13977778888')[1]['values']['archive_no'],source['archive_no'])
        row=next(row for row in self.call('employees')[1] if row['id']==employee)
        self.assertEqual(self.call('employees/'+str(employee),dict(row,archive_no='00000000000000'),method='PUT')[0],200)
        self.assertEqual(next(row for row in self.call('employees')[1] if row['id']==employee)['archive_no'],source['archive_no'])

    def test_25_pdf_library_candidate_import(self):
        pdf={'name':'简历.pdf','content':base64.b64encode(b'%PDF-1.4\nresume-library').decode()}
        self.assertEqual(self.call('resume_documents',{'name':'缺少文件'})[0],400)
        status,created=self.call('resume_documents',{'name':'识别测试','_resume_upload':pdf})
        self.assertEqual(status,201)
        document=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        viewer=self.call('me')[1]
        self.assertEqual(document['created_by'],viewer['nickname'] or viewer['username'])
        self.assertEqual(document['resume_name'],'简历.pdf')
        self.assertEqual(document['name'],'简历.pdf')
        self.assertEqual(document['source_id'],document['archive_no'])
        self.assertEqual(len(document['file_hash']),64)
        self.assertEqual(document['status'],'待识别')
        self.assertTrue(any(row['id']==created['id'] for row in self.call('resume_documents/pending')[1]))
        self.assertEqual(self.call('resume_documents',{'name':'重复上传','_resume_upload':pdf})[0],400)
        self.assertEqual(self.call('resume_documents',public=True)[0],401)
        self.assertEqual(self.call('dify-embed',public=True)[0],401)
        with patch.dict('os.environ',{},clear=True):
            self.assertEqual(self.call('dify-embed')[1]['url'],'http://biaozhun.njncc.com/workflow/e3JaA7Tg8GrWKJi5')
        path='resume_documents/'+str(created['id'])+'/candidate'
        self.assertEqual(self.call(path,{'name':'缺字段'})[0],400)
        self.assertEqual(next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])['status'],'待识别')
        self.assertTrue(any(row['id']==created['id'] for row in self.call('resume_documents/pending')[1]))
        self.assertEqual(self.call('resume_documents/'+str(created['id']),{'status':'已识别','source_id':'伪造','created_by':'伪造创建人'},method='PUT')[0],200)
        self.assertEqual(next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])['source_id'],document['source_id'])
        self.assertEqual(next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])['created_by'],document['created_by'])
        status,result=self.call(path,{'name':'识别人才','phone':'13945671238','education':'本科','school_name':'识别大学','major':'土木工程'})
        self.assertEqual(status,201)
        candidate=next(row for row in self.call('candidates')[1] if row['id']==result['id'])
        self.assertEqual(candidate['resume'],document['resume'])
        self.assertEqual(next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])['status'],'已入库')
        self.assertEqual(self.call('resume_documents/'+str(created['id']),{'status':'待识别'},method='PUT')[0],200)
        self.assertFalse(any(row['id']==created['id'] for row in self.call('resume_documents/pending')[1]))
        self.assertEqual(self.call('resume_documents/pending',public=True)[0],401)
        self.assertEqual(self.call(path,{'name':'重复人才','phone':'13945671238','education':'本科','school_name':'识别大学','major':'土木工程'})[1]['id'],result['id'])
        self.assertEqual(self.call('resume_documents/'+str(created['id']),method='DELETE')[0],200)
        self.assertEqual(self.call('attachments/'+document['resume'])[0],200)

    def test_36_deleted_candidate_releases_resume_for_recognition_and_reimport(self):
        pdf={'name':'删除后重识别.pdf','content':base64.b64encode(b'%PDF-1.4\nrecognize-after-delete').decode()}
        created=self.call('resume_documents',{'name':'删除后重识别','_resume_upload':pdf})[1]
        document=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        path='resume_documents/'+str(created['id'])+'/candidate'
        first=self.call(path,{'name':'首次人才','phone':'13987651001','education':'本科','school_name':'测试大学','major':'工程'})[1]

        preview=self.call('candidates/delete-preview',{'ids':[first['id']]})[1]
        self.assertEqual(self.call('candidates/batch-delete',{'ids':[first['id']],'expected':preview['items']})[0],200)

        source=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        self.assertEqual(source['status'],'待识别')
        self.assertTrue(any(row['id']==created['id'] for row in self.call('resume_documents/pending')[1]))
        self.assertEqual(self.call('attachments/'+document['resume'])[0],200)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM resume_imports WHERE source_id IN (?,?)',(document['source_id'],document['file_hash'])).fetchone()[0],0)

        status,second=self.call(path,{'name':'重新识别人才','phone':'13987651002','education':'本科','school_name':'测试大学','major':'工程'})
        self.assertEqual(status,201)
        self.assertNotEqual(second['id'],first['id'])
        self.assertEqual(self.call(path,{'name':'重复人才','phone':'13987651003','education':'本科','school_name':'测试大学','major':'工程'})[1]['id'],second['id'])

    def test_37_startup_releases_historical_import_whose_candidate_was_deleted(self):
        pdf={'name':'历史删除记录.pdf','content':base64.b64encode(b'%PDF-1.4\nhistorical-deleted-candidate').decode()}
        created=self.call('resume_documents',{'name':'历史删除记录','_resume_upload':pdf})[1]
        document=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        candidate=self.call('resume_documents/'+str(created['id'])+'/candidate',{'name':'历史人才','phone':'13987651004','education':'本科','school_name':'测试大学','major':'工程'})[1]
        with server.db() as conn:
            conn.execute('PRAGMA foreign_keys=OFF')
            conn.execute('DELETE FROM candidates WHERE id=?',(candidate['id'],))
            conn.execute("UPDATE resume_documents SET status='已识别' WHERE id=?",(created['id'],))
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM resume_imports WHERE candidate_id=? AND source_id IN (?,?)',(candidate['id'],document['source_id'],document['file_hash'])).fetchone()[0],2)

        server.initialize()

        source=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        self.assertEqual(source['status'],'待识别')
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM resume_imports WHERE source_id IN (?,?)',(document['source_id'],document['file_hash'])).fetchone()[0],0)

    def test_38_orphaned_shared_hash_does_not_release_a_live_candidate_document(self):
        first=self.call('resume_documents',{'name':'历史重复一','_resume_upload':{'name':'历史重复一.pdf','content':base64.b64encode(b'%PDF-1.4\nhistorical-one').decode()}})[1]
        second=self.call('resume_documents',{'name':'历史重复二','_resume_upload':{'name':'历史重复二.pdf','content':base64.b64encode(b'%PDF-1.4\nhistorical-two').decode()}})[1]
        first_document=next(row for row in self.call('resume_documents')[1] if row['id']==first['id'])
        second_document=next(row for row in self.call('resume_documents')[1] if row['id']==second['id'])
        first_candidate=self.call('resume_documents/'+str(first['id'])+'/candidate',{'name':'历史重复人才一','phone':'13987651005','education':'本科','school_name':'测试大学','major':'工程'})[1]
        second_candidate=self.call('resume_documents/'+str(second['id'])+'/candidate',{'name':'历史重复人才二','phone':'13987651006','education':'本科','school_name':'测试大学','major':'工程'})[1]
        with server.db() as conn:
            conn.execute('UPDATE resume_documents SET file_hash=? WHERE id=?',(first_document['file_hash'],second['id']))
            conn.execute('DELETE FROM candidates WHERE id=?',(first_candidate['id'],))

        server.initialize()

        source=next(row for row in self.call('resume_documents')[1] if row['id']==second['id'])
        self.assertEqual(source['status'],'已入库')
        with server.db() as conn:
            marker=conn.execute('SELECT candidate_id FROM resume_imports WHERE source_id=?',(first_document['file_hash'],)).fetchone()
            self.assertEqual(marker['candidate_id'],second_candidate['id'])
            self.assertEqual(conn.execute('SELECT candidate_id FROM resume_imports WHERE source_id=?',(second_document['source_id'],)).fetchone()['candidate_id'],second_candidate['id'])

    def test_39_recognition_cannot_overwrite_pending_state_after_import_delete_cycle(self):
        created=self.call('resume_documents',{'name':'识别删除竞态','recruitment_type':'校园招聘','channel':'校园平台','_resume_upload':{'name':'识别删除竞态.pdf','content':base64.b64encode(b'%PDF-1.4\nrecognition-delete-race').decode()}})[1]
        source=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        def import_delete_cycle(*_):
            with server.db() as conn:
                conn.execute("UPDATE resume_documents SET status='已入库' WHERE id=?",(created['id'],))
                conn.execute("UPDATE resume_documents SET status='待识别',updated_at='2099-01-01T00:00:00.123456' WHERE id=?",(created['id'],))
            return {'candidate':{'basic':{'name':'不应覆盖'}}}
        with patch('server.dify_client.configuration',return_value={}),patch('server.dify_client.recognize',side_effect=import_delete_cycle):
            status,_=self.call('resume_documents/'+str(created['id'])+'/recognize',{})
        self.assertEqual(status,409)
        current=next(row for row in self.call('resume_documents')[1] if row['id']==created['id'])
        self.assertEqual(current['status'],'待识别')
        self.assertNotEqual(current['updated_at'],source['updated_at'])

    def test_24_attachment_prefill_and_sync(self):
        identity='11010519491231002X'
        source=self.call('applications',{'name':'附件联动','school_name':'测试大学','identity_card':identity,'_resume_upload':{'name':'联动.pdf','content':base64.b64encode(b'%PDF-1.4\nlinked').decode()}})[1]['id']
        response=self.call('interview-prefill?identity_card='+identity)[1]
        attachment=response['attachment_fields']['resume']
        self.assertEqual(attachment['name'],'联动.pdf')
        self.assertEqual(self.call('attachments/'+attachment['id']+'?preview=1')[0],200)
        identifier=self.call('interviews',{'identity_card':identity,'_remove_resume':True})[1]['id']
        row=next(row for row in self.call('interviews')[1] if row['id']==identifier)
        self.assertEqual(row['resume'],'')
        self.assertEqual(self.call('interviews/'+str(identifier),dict(row,_sync_attachments=True),method='PUT')[0],200)
        row=next(row for row in self.call('interviews')[1] if row['id']==identifier)
        self.assertEqual(row['resume'],attachment['id'])
        self.assertEqual(self.call('interviews/'+str(identifier),dict(row,_sync_attachments=True,_remove_resume=True),method='PUT')[0],200)
        self.assertEqual(next(row for row in self.call('interviews')[1] if row['id']==identifier)['resume'],'')

    def test_23_selective_cascade_preserves_records(self):
        candidate=self.call('candidates',{'name':'保留关联测试','phone':'13956781234','school_name':'测试大学','major':'工程','education':'本科','_resume_upload':{'name':'保留.pdf','content':base64.b64encode(b'%PDF-1.4\nkeep').decode()}})[1]['id']
        application=self.call('applications',{'phone':'13956781234'})[1]['id']
        original=next(row for row in self.call('applications')[1] if row['id']==application)
        preview=self.call('candidates/delete-preview',{'ids':[candidate]})[1]
        payload={'ids':[candidate],'expected':preview['items'],'selected':[]}
        self.assertEqual(self.call('candidates/batch-delete',payload)[0],400)
        payload['selected']=[{'entity':'candidates','id':candidate},{'entity':'employees','id':999999}]
        self.assertEqual(self.call('candidates/batch-delete',payload)[0],400)
        payload['selected']=[{'entity':'candidates','id':candidate}]
        status,result=self.call('candidates/batch-delete',payload)
        self.assertEqual((status,result['count']),(200,1))
        kept=next(row for row in self.call('applications')[1] if row['id']==application)
        self.assertIsNone(kept['candidate_id'])
        self.assertEqual(kept['name'],original['name'])
        self.assertEqual(kept['resume'],original['resume'])
        self.assertEqual(self.call('attachments/'+kept['resume'])[0],200)
        status,content=self.call('attachments/'+kept['resume']+'?preview=1')
        self.assertEqual(status,200)
        self.assertTrue(content.startswith(b'%PDF'))
        self.assertEqual(self.call('attachments/'+kept['resume']+'?preview=1',public=True)[0],401)

    def test_21_member_accounts_and_permissions(self):
        self.assertEqual(self.call('members',public=True)[0],401)
        self.assertEqual(self.call('members',{'username':'admin','password':'member-pass-123'})[0],400)
        self.assertEqual(self.call('members',{'username':'recruiter','password':'short'})[0],400)
        self.assertEqual(self.call('members',{'username':'recruiter','password':'member-pass-123','phone':'invalid'})[0],400)
        self.assertEqual(self.call('members',{'username':'nickname-too-long','password':'member-pass-123','nickname':'长'*31})[0],400)
        self.assertEqual(self.call('members',{'username':'nickname-control','password':'member-pass-123','nickname':'换\n行'})[0],400)
        self.assertEqual(self.call('members',{'username':'recruiter','password':'member-pass-123','role':'admin','phone':'13912345678','nickname':'初始昵称'})[0],201)
        self.assertEqual(self.call('members',{'username':'RECRUITER','password':'member-pass-123'})[0],409)
        member=build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))
        def request(path,data=None,method=None):
            req=Request(self.url+'/api/'+path,data=json.dumps(data).encode() if data is not None else None,headers={'Content-Type':'application/json'},method=method)
            try: response=member.open(req)
            except HTTPError as error: response=error
            with response: return response.status,json.loads(response.read())
        self.assertEqual(request('login',{'username':'recruiter','password':'wrong'})[0],401)
        self.assertEqual(request('login',{'username':'RECRUITER','password':'member-pass-123'})[0],200)
        self.assertEqual(request('me')[1]['role'],'member')
        self.assertEqual(request('me')[1]['nickname'],'初始昵称')
        admin_before=self.call('me')[1]
        self.assertEqual(request('me')[1]['phone'],'13912345678')
        self.assertEqual(request('me/phone',{'phone':'13887654321','id':admin_before['id'],'role':'admin'},method='PUT')[0],200)
        self.assertEqual(request('me')[1]['phone'],'13887654321')
        self.assertEqual(self.call('me')[1],admin_before)
        self.assertEqual(request('me/phone',{'phone':'invalid'},method='PUT')[0],400)
        self.assertEqual(self.call('me/phone',{'phone':'13887654321'},method='PUT',public=True)[0],401)
        self.assertEqual(request('me',{'nickname':'招聘小李','id':admin_before['id'],'role':'admin'},method='PUT')[0],200)
        self.assertEqual(request('me')[1]['nickname'],'招聘小李')
        self.assertEqual(request('me')[1]['role'],'member')
        self.assertEqual(self.call('me')[1],admin_before)
        self.assertEqual(self.call('members')[1][0]['nickname'],'招聘小李')
        self.assertEqual(request('me',{'nickname':'长'*31},method='PUT')[0],400)
        self.assertEqual(request('me',{'nickname':'换\n行'},method='PUT')[0],400)
        self.assertEqual(self.call('me',{'nickname':'招聘管理员'},method='PUT')[0],200)
        self.assertEqual(self.call('me')[1]['nickname'],'招聘管理员')
        self.assertEqual(self.call('me/phone',{'phone':'025-12345678'},method='PUT')[0],200)
        self.assertEqual(self.call('me')[1]['phone'],'025-12345678')
        self.assertEqual(self.call('me',{'nickname':'未登录'},method='PUT',public=True)[0],401)
        self.assertEqual(request('members')[0],200)
        self.assertNotIn('password',request('members')[1][0])
        self.assertNotIn('salt',request('members')[1][0])
        self.assertEqual(request('members',{'username':'another','password':'member-pass-123'})[0],403)
        fixtures={'schools':{'name':'成员院校','province':'江苏省','city':'南京市'},'candidates':{'name':'成员人才','phone':'13933334444','school_name':'成员院校','education':'本科','major':'工程'},'applications':{'name':'成员应聘','school_name':'成员院校'},'interviews':{'name':'成员面试','school_name':'成员院校'},'employees':{'name':'成员入职'},'questions':{'question':'成员问题'}}
        for entity,payload in fixtures.items():
            status,result=request(entity,payload)
            self.assertEqual(status,201,(entity,result))
            identifier=result['id']
            self.assertTrue(any(row['id']==identifier for row in request(entity)[1]))
            self.assertEqual(request(entity+'/'+str(identifier),payload,method='PUT')[0],200)
            self.assertEqual(request(entity+'/'+str(identifier),method='DELETE')[0],200)
        server.initialize()
        self.assertEqual(request('me')[1]['nickname'],'招聘小李')
        with server.db() as conn:
            saved=conn.execute('SELECT * FROM members WHERE username=?',('recruiter',)).fetchone()
            self.assertNotEqual(saved['password'],'member-pass-123')
        self.assertNotIn('password',self.call('members')[1][0])
        member_id=request('me')[1]['id']
        role_path='members/'+str(member_id)+'/role'
        self.assertEqual(request(role_path,{'role':'manager'},method='PUT')[0],403)
        self.assertEqual(self.call(role_path,{'role':'manager'},method='PUT',public=True)[0],401)
        self.assertEqual(self.call(role_path,{'role':'admin'},method='PUT')[0],400)
        self.assertEqual(self.call(role_path,{'role':'manager'},method='PUT')[0],200)
        self.assertEqual(request('me')[1]['role'],'manager')
        self.assertEqual(request(role_path,{'role':'member'},method='PUT')[0],403)
        self.assertEqual(request('members',{'username':'managed-child','password':'child-password','role':'manager','is_admin':1})[0],201)
        child=next(row for row in request('members')[1] if row['username']=='managed-child')
        self.assertEqual(child['role'],'member')
        self.assertEqual(request('members/'+str(child['id'])+'/role',{'role':'manager'},method='PUT')[0],403)
        self.assertEqual(request('members/'+str(child['id']),method='DELETE')[0],200)
        self.assertEqual(request('members/admin',method='DELETE')[0],400)
        server.initialize()
        self.assertEqual(request('me')[1]['role'],'manager')
        self.assertEqual(request('password',{'old_password':'wrong','password':'new-member-pass'})[0],400)
        self.assertEqual(request('password',{'old_password':'member-pass-123','password':'new-member-pass'})[0],200)
        self.assertEqual(request('schema')[0],401)
        self.assertEqual(self.call('me')[1]['role'],'admin')
        self.assertEqual(request('login',{'username':'recruiter','password':'member-pass-123'})[0],401)
        self.assertEqual(request('login',{'username':'recruiter','password':'new-member-pass'})[0],200)
        self.assertEqual(request('me')[1]['role'],'manager')
        self.assertEqual(self.call(role_path,{'role':'member'},method='PUT')[0],200)
        self.assertEqual(request('me')[1]['role'],'member')
        self.assertEqual(request('members',{'username':'forbidden-child','password':'child-password'})[0],403)
        self.assertEqual(request('me')[1]['phone'],'13887654321')
        self.assertEqual(request('me/phone',{'phone':''},method='PUT')[0],200)
        self.assertEqual(request('me')[1]['phone'],'')
        self.assertEqual(request('me')[1]['nickname'],'招聘小李')
        self.assertEqual(request('me',{'nickname':''},method='PUT')[0],200)
        self.assertEqual(request('me')[1]['nickname'],'')
        member_id=request('me')[1]['id']
        counts=self.call('stats')[1]['counts']
        self.assertEqual(request('members/'+str(member_id),method='DELETE')[0],403)
        self.assertEqual(self.call('members/'+str(member_id),method='DELETE',public=True)[0],401)
        self.assertEqual(self.call('members/'+str(member_id),method='DELETE')[0],200)
        self.assertEqual(request('schema')[0],401)
        self.assertEqual(request('login',{'username':'recruiter','password':'new-member-pass'})[0],401)
        self.assertEqual(self.call('stats')[1]['counts'],counts)
        self.assertEqual(self.call('members/'+str(member_id),method='DELETE')[0],404)
        self.assertEqual(self.call('members',{'username':'recruiter','password':'replacement-pass'})[0],201)
        self.assertEqual(request('schema')[0],401)
        self.assertEqual(request('login',{'username':'recruiter','password':'replacement-pass'})[0],200)
        self.assertEqual(request('logout',{})[0],200)
        self.assertEqual(request('schema')[0],401)

    def test_22_primary_admin_can_reset_member_password_and_revoke_sessions(self):
        def client_request(client, path, data=None, method=None):
            request = Request(
                self.url + '/api/' + path,
                data=json.dumps(data).encode() if data is not None else None,
                headers={'Content-Type': 'application/json'},
                method=method,
            )
            try:
                response = client.open(request)
            except HTTPError as error:
                response = error
            with response:
                return response.status, json.loads(response.read())

        self.assertEqual(self.call('members', {
            'username': 'reset-target', 'password': 'old-target-password',
        })[0], 201)
        self.assertEqual(self.call('members', {
            'username': 'reset-manager', 'password': 'manager-password',
        })[0], 201)
        members = {member['username']: member for member in self.call('members')[1]}
        target_id = members['reset-target']['id']
        manager_id = members['reset-manager']['id']
        self.assertEqual(self.call(
            f'members/{manager_id}/role', {'role': 'manager'}, method='PUT'
        )[0], 200)

        target = build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))
        manager = build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))
        self.assertEqual(client_request(target, 'login', {
            'username': 'reset-target', 'password': 'old-target-password',
        })[0], 200)
        self.assertEqual(client_request(manager, 'login', {
            'username': 'reset-manager', 'password': 'manager-password',
        })[0], 200)
        path = f'members/{target_id}/password'

        self.assertEqual(self.call(
            path, {'password': 'new-target-password'}, public=True, method='PUT'
        )[0], 401)
        self.assertEqual(client_request(
            manager, path, {'password': 'new-target-password'}, method='PUT'
        )[0], 403)
        self.assertEqual(self.call(path, {'password': 'short'}, method='PUT')[0], 400)
        self.assertEqual(self.call(
            path, {'password': 'new-target-password'}, method='PUT'
        )[0], 200)

        self.assertEqual(client_request(target, 'schema')[0], 401)
        self.assertEqual(client_request(target, 'login', {
            'username': 'reset-target', 'password': 'old-target-password',
        })[0], 401)
        self.assertEqual(client_request(target, 'login', {
            'username': 'reset-target', 'password': 'new-target-password',
        })[0], 200)
        self.assertEqual(client_request(manager, 'schema')[0], 200)
